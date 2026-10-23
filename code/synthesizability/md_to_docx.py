# -*- coding: utf-8 -*-
"""
Convert the manuscript markdown to a submission-ready .docx.

Handles exactly what MANUSCRIPT.md uses: ATX headings, paragraphs with **bold**/*italic*/
`code` runs, pipe tables, images written as {{artifact:art_<id>}}, bullet lists,
horizontal rules and blockquotes. Figures are resolved through host.artifact_path and
embedded at 16 cm width; the paragraph immediately following an image is styled as a caption.

Double-spaced, 12 pt serif, continuous line numbers - the format most journals ask for
at submission.
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

IMG_RE   = re.compile(r"^!\[(?P<alt>[^\]]*)\]\(\{\{artifact:art_(?P<aid>[0-9a-f-]+)\}\}\)\s*$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_line_numbers(section):
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sectPr.append(ln)


def _runs(par, text):
    """Add text to a paragraph, honouring **bold**, *italic* and `code`."""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10.5)
        elif tok.startswith("*") and tok.endswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def _table(doc, rows):
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = []
    for r in rows[2:]:
        body.append([c.strip() for c in r.strip().strip("|").split("|")])
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _runs(cell.paragraphs[0], h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for brow in body:
        cells = t.add_row().cells
        for i, v in enumerate(brow[:len(header)]):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], v)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    return t


def convert(md_path, out_path, resolve_artifact=None, double_space=True):
    """resolve_artifact: callable(artifact_id) -> local image path, or None to skip images."""
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0)
    if double_space:
        st.paragraph_format.line_spacing = 2.0
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        _add_line_numbers(s)

    i, after_image, n_img, missing = 0, False, 0, []
    while i < len(lines):
        ln = lines[i]

        m = IMG_RE.match(ln)
        if m:
            path = resolve_artifact(m.group("aid")) if resolve_artifact else None
            if path and os.path.exists(path):
                doc.add_picture(path, width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                n_img += 1
            else:
                missing.append(m.group("alt") or m.group("aid"))
                p = doc.add_paragraph()
                r = p.add_run(f"[figure not embedded: {m.group('alt')}]")
                r.italic = True; r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            after_image = True
            i += 1
            continue

        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("#").strip(), level=min(lvl, 4))
            after_image = False
            i += 1
            continue

        if ln.strip() in ("---", "***", "___"):
            i += 1
            continue

        if ln.lstrip().startswith("|") and i + 1 < len(lines) and set(
                lines[i + 1].replace("|", "").replace(":", "").strip()) <= {"-", " "}:
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i]); i += 1
            _table(doc, block)
            doc.add_paragraph()
            after_image = False
            continue

        if ln.lstrip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, ln.lstrip()[2:])
            after_image = False
            i += 1
            continue

        if not ln.strip():
            i += 1
            continue

        # gather a paragraph (markdown soft-wraps inside one)
        buf = []
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") \
                and not lines[i].lstrip().startswith(("|", "- ", "* ")) \
                and not IMG_RE.match(lines[i]) and lines[i].strip() not in ("---", "***", "___"):
            buf.append(lines[i].strip()); i += 1
        text = " ".join(buf)

        p = doc.add_paragraph()
        _runs(p, text.lstrip("> "))
        if after_image:
            for r in p.runs:
                r.font.size = Pt(10)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(14)
            after_image = False
        else:
            p.paragraph_format.space_after = Pt(6)

    doc.save(out_path)
    return dict(images_embedded=n_img, images_missing=missing,
                paragraphs=len(doc.paragraphs), tables=len(doc.tables))
